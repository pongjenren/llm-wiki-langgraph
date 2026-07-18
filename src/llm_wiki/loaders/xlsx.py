"""Excel loader: renders each sheet as a markdown table.

Cells are read with data_only=True so formulas come through as their last
cached value rather than as formula text.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Guard against a runaway sheet producing a document too large to reason about.
MAX_ROWS_PER_SHEET = 500


def _fmt(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return str(value).replace("|", r"\|").replace("\n", " ")


def load_xlsx(path: Path) -> str:
    workbook = load_workbook(path, data_only=True, read_only=True)
    parts: list[str] = []

    for sheet in workbook.worksheets:
        rows = [
            [_fmt(cell) for cell in row]
            for row in sheet.iter_rows(max_row=MAX_ROWS_PER_SHEET, values_only=True)
        ]
        # Drop wholly empty trailing rows openpyxl often reports for sparse sheets.
        while rows and not any(cell for cell in rows[-1]):
            rows.pop()

        if not rows:
            parts.append(f"## Sheet: {sheet.title}\n\n_(empty)_")
            continue

        header, *body = rows
        width = max(len(r) for r in rows)
        header = header + [""] * (width - len(header))

        # Table rows must stay on consecutive lines or the markdown table breaks,
        # so the whole sheet is assembled as one block.
        lines = [
            f"## Sheet: {sheet.title}",
            "",
            "| " + " | ".join(header) + " |",
            "| " + " | ".join("---" for _ in header) + " |",
        ]
        lines += ["| " + " | ".join(row + [""] * (width - len(row))) + " |" for row in body]

        if sheet.max_row and sheet.max_row > MAX_ROWS_PER_SHEET:
            lines.append("")
            lines.append(f"_(truncated: showing first {MAX_ROWS_PER_SHEET} of {sheet.max_row} rows)_")

        parts.append("\n".join(lines))

    workbook.close()
    return "\n\n".join(parts)
